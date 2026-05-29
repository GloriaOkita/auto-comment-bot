#!/usr/bin/env python3
"""
自动评论机 — 图形界面
"""

import json
import os
import threading
from datetime import datetime
from pathlib import Path
from tkinter import (
    END,
    LEFT,
    RIGHT,
    Button,
    Entry,
    filedialog,
    Frame,
    Label,
    LabelFrame,
    messagebox,
    OptionMenu,
    Scrollbar,
    StringVar,
    Text,
    Tk,
)

BASE_DIR = Path(__file__).parent
CONFIG_PATH = BASE_DIR / "config.json"


class App:
    def __init__(self):
        self.root = Tk()
        self.root.title("自动评论机")
        self.root.geometry("680x600")
        self.root.minsize(500, 420)
        self.root.configure(bg="#f5f5f5")

        self.bot = None
        self.bot_thread = None
        self.running = False
        self._paused = False

        self._load_config()
        self._build_ui()
        self._refresh_file_status()

    def _load_config(self):
        self.config = {}
        if CONFIG_PATH.exists():
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                self.config = json.load(f)

    def _save_config(self):
        self.config["last_accounts_file"] = self.accounts_path
        self.config["last_post_file"] = self.posts_path
        self.config["api_key"] = self.apikey_var.get()
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(self.config, f, ensure_ascii=False, indent=2)

    # ============================================================
    #  UI
    # ============================================================
    def _build_ui(self):
        # ---------- 标题 ----------
        Label(
            self.root, text="自动评论机",
            font=("", 20, "bold"), bg="#f5f5f5", fg="#333",
        ).pack(pady=(16, 4))
        Label(
            self.root, text="自动登录 · AI生成评论 · 发布 · 标记完成",
            bg="#f5f5f5", fg="#888", font=("", 10),
        ).pack(pady=(0, 12))

        # ---------- 文件上传区 ----------
        file_frame = LabelFrame(
            self.root, text="上传文件", padx=14, pady=12,
            font=("", 11, "bold"),
        )
        file_frame.pack(fill="x", padx=20, pady=(0, 10))

        # 账号文件行
        self.accounts_path = self.config.get("last_accounts_file", "")
        self._build_upload_row(
            file_frame, "上传账号文件",
            "accounts_path", "accounts_status_label",
            self._pick_accounts,
        )

        # 分隔
        Frame(file_frame, height=6).pack()

        # 帖子文件行
        default_post = self.config.get("last_post_file", "捣谷社区每日打标帖.xlsx")
        if not os.path.isabs(default_post):
            default_post = str(BASE_DIR / default_post)
        self.posts_path = default_post

        self._build_upload_row(
            file_frame, "上传评论表格",
            "posts_path", "posts_status_label",
            self._pick_posts,
        )

        # ---------- API Key ----------
        api_frame = Frame(self.root, bg="#f5f5f5")
        api_frame.pack(fill="x", padx=20, pady=(0, 6))
        Label(api_frame, text="API Key", bg="#f5f5f5", font=("", 10)).pack(
            side=LEFT, padx=(0, 8)
        )
        self.apikey_var = StringVar(value=self.config.get("api_key", ""))
        Entry(
            api_frame, textvariable=self.apikey_var, width=52,
            font=("", 10),
        ).pack(side=LEFT)

        # ---------- 员工选择 ----------
        emp_frame = Frame(self.root, bg="#f5f5f5")
        emp_frame.pack(fill="x", padx=20, pady=(0, 8))
        Label(
            emp_frame, text="评论谁的内容？", bg="#f5f5f5", font=("", 10),
        ).pack(side=LEFT, padx=(0, 8))
        self.emp_var = StringVar(value="全部")
        self.emp_menu = OptionMenu(emp_frame, self.emp_var, "全部")
        self.emp_menu.pack(side=LEFT)

        # ---------- 启动/进度按钮 ----------
        self.action_btn = Button(
            self.root,
            text="开始评论",
            font=("", 14, "bold"),
            relief="flat",
            padx=30,
            pady=8,
            command=self._toggle_run,
        )
        self.action_btn.pack(pady=(4, 8))

        # ---------- 日志 ----------
        log_frame = LabelFrame(
            self.root, text="运行日志", padx=10, pady=6,
            font=("", 10, "bold"),
        )
        log_frame.pack(fill="both", expand=True, padx=20, pady=(0, 14))

        self.log_text = Text(
            log_frame,
            bg="#1e1e1e", fg="#d4d4d4",
            insertbackground="white",
            wrap="word",
            font=("Menlo", 10),
            relief="flat", borderwidth=0,
        )
        scrollbar = Scrollbar(log_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.pack(side=LEFT, fill="both", expand=True)
        scrollbar.pack(side=RIGHT, fill="y")

        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self._update_emp_menu()

    def _build_upload_row(self, parent, btn_text, path_attr, status_attr, pick_cmd):
        """构建一行：按钮（状态直接写在按钮文字里）"""
        row = Frame(parent)
        row.pack(fill="x", pady=3)

        btn = Button(
            row, text=btn_text, font=("", 10),
            width=28, anchor="w", padx=8,
            command=pick_cmd,
        )
        btn.pack(side=LEFT, padx=(0, 12))

        setattr(self, status_attr, btn)

    # ============================================================
    #  交互
    # ============================================================
    def _pick_accounts(self):
        path = filedialog.askopenfilename(
            title="选择账号 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.accounts_path = path
            self._refresh_file_status()
            self._update_emp_menu()

    def _pick_posts(self):
        path = filedialog.askopenfilename(
            title="选择帖子 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.posts_path = path
            self._refresh_file_status()

    def _refresh_file_status(self):
        """刷新文件状态文字"""
        # 账号文件
        acc_path = getattr(self, "accounts_path", "")
        if acc_path and os.path.exists(acc_path):
            try:
                from comment_bot import AccountManager
                mgr = AccountManager(acc_path)
                names = mgr.get_employee_names()
                total = sum(len(mgr.get_accounts(n)) for n in names)
                if total > 0:
                    self.accounts_status_label.configure(
                        text=f"上传账号文件  ✓ 已上传（{total}个账号）"
                    )
                else:
                    self.accounts_status_label.configure(
                        text="上传账号文件  ⚠ 无账号数据"
                    )
            except Exception as e:
                self.accounts_status_label.configure(text=f"上传账号文件  ✗ 读取失败")
        else:
            self.accounts_status_label.configure(text="上传账号文件")

        # 帖子文件
        post_path = getattr(self, "posts_path", "")
        if post_path and os.path.exists(post_path):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(post_path, data_only=True)
                sheet_count = len(wb.sheetnames)
                total_rows = sum(wb[s].max_row - 1 for s in wb.sheetnames)
                wb.close()
                self.posts_status_label.configure(
                    text=f"上传评论表格  ✓ 已上传（{sheet_count}个表，{total_rows}条）"
                )
            except Exception as e:
                self.posts_status_label.configure(text=f"上传评论表格  ✗ 读取失败")
        else:
            self.posts_status_label.configure(text="上传评论表格")

    def _update_emp_menu(self):
        menu = self.emp_menu["menu"]
        menu.delete(0, "end")
        menu.add_command(label="全部", command=lambda: self.emp_var.set("全部"))
        acc_path = getattr(self, "accounts_path", "")
        if acc_path and os.path.exists(acc_path):
            try:
                from comment_bot import AccountManager
                mgr = AccountManager(acc_path)
                for name in mgr.get_employee_names():
                    menu.add_command(
                        label=name, command=lambda n=name: self.emp_var.set(n)
                    )
            except Exception:
                pass

    # ============================================================
    #  运行控制
    # ============================================================
    def _log(self, msg):
        ts = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(END, f"[{ts}] {msg}\n")
        self.log_text.see(END)

    def _update_progress(self, msg):
        """从后台线程更新按钮文字"""
        self.root.after(0, lambda: self.action_btn.configure(text=msg))

    def _toggle_run(self):
        if self.running:
            # 暂停/继续
            self._paused = not self._paused
            if self._paused:
                self.bot.cancel()
                self.action_btn.configure(text="继续评论")
            else:
                self.action_btn.configure(text="暂停中...")
                # 重新启动
                self._start_bot_thread()
        else:
            self._start()

    def _start(self):
        acc_path = getattr(self, "accounts_path", "")
        if not acc_path or not os.path.exists(acc_path):
            messagebox.showerror("错误", "请先上传账号文件")
            return

        post_path = getattr(self, "posts_path", "")
        if not post_path or not os.path.exists(post_path):
            messagebox.showerror("错误", "请先上传评论表格")
            return

        self._save_config()
        self.log_text.delete("1.0", END)
        self._log("自动评论机启动")

        self.running = True
        self._paused = False
        self.action_btn.configure(text="暂停")

        self._start_bot_thread()

    def _start_bot_thread(self):
        emp_filter = self.emp_var.get()
        if emp_filter == "全部":
            emp_filter = None

        self.bot_thread = threading.Thread(
            target=self._run_bot,
            args=(self.accounts_path, self.posts_path, emp_filter),
            daemon=True,
        )
        self.bot_thread.start()

    def _run_bot(self, accounts_path, posts_path, employee_filter):
        from comment_bot import CommentBot, ConfigManager

        config = ConfigManager()
        config.data.update(self.config)
        config.data["last_post_file"] = posts_path

        def _on_progress(msg):
            self.root.after(0, lambda: self.action_btn.configure(text=msg))

        self.bot = CommentBot(
            config, log_callback=self._log, progress_callback=_on_progress,
        )
        try:
            self.bot.run(accounts_path, employee_filter)
        except Exception as e:
            import traceback
            self._log(f"运行错误: {e}")
            self._log(traceback.format_exc())
        finally:
            self.root.after(0, self._on_finish)

    def _on_finish(self):
        self.running = False
        self._paused = False
        self.action_btn.configure(text="开始评论")

    def _on_close(self):
        if self.running:
            ok = messagebox.askokcancel("确认", "正在运行中，确定退出吗？")
            if not ok:
                return
            if self.bot:
                self.bot.cancel()
        self._save_config()
        self.root.destroy()


def main():
    app = App()
    app.root.mainloop()


if __name__ == "__main__":
    main()
